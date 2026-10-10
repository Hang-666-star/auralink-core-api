#!/usr/bin/env python3
"""Validate a painting workbook and image directory before database import."""

from __future__ import annotations

import argparse
import hashlib
import json
import sys
import unicodedata
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EXPECTED_HEADERS = [
    "入库编号",
    "源作品序号",
    "作者姓名",
    "作品名称",
    "该作品参加过哪些展览",
    "创作年份",
    "创作归属/参与方式",
    "主题材",
    "题材2",
    "创作风格",
    "作品材料与载体",
    "形制",
    "尺寸",
    "宽(cm)",
    "高(cm)",
    "所属系列",
    "收藏状态",
    "高清原图文件名",
    "备注",
]
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp", ".heic"}


def clean_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = unicodedata.normalize("NFC", str(value)).strip()
    return text or None


def normalized_name(value: Any) -> str | None:
    text = clean_text(value)
    return text.casefold() if text else None


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _issue(severity: str, code: str, message: str, **details: Any) -> dict[str, Any]:
    return {"severity": severity, "code": code, "message": message, "details": details}


def read_workbook(path: Path) -> tuple[str, list[dict[str, Any]], list[dict[str, Any]]]:
    issues: list[dict[str, Any]] = []
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(iterator)
        except StopIteration:
            return sheet.title, [], [_issue("error", "empty_workbook", "主表没有表头")]

        headers = [clean_text(value) for value in raw_headers]
        missing = [header for header in EXPECTED_HEADERS if header not in headers]
        if missing:
            issues.append(_issue("error", "missing_headers", "主表缺少必填字段", headers=missing))

        duplicate_headers = sorted({header for header in headers if header and headers.count(header) > 1})
        if duplicate_headers:
            issues.append(_issue("error", "duplicate_headers", "主表存在重复字段", headers=duplicate_headers))

        rows: list[dict[str, Any]] = []
        for row_number, values in enumerate(iterator, start=2):
            if not any(clean_text(value) for value in values):
                continue
            record = {header: values[index] if index < len(values) else None for index, header in enumerate(headers) if header}
            record["__row_number"] = row_number
            rows.append(record)
        return sheet.title, rows, issues
    finally:
        workbook.close()


def validate_batch(
    workbook_path: Path,
    image_directory: Path,
    expected_count: int | None = None,
    hash_images: bool = True,
) -> dict[str, Any]:
    workbook_path = workbook_path.resolve()
    image_directory = image_directory.resolve()
    issues: list[dict[str, Any]] = []

    if not workbook_path.is_file():
        raise FileNotFoundError(f"主表不存在: {workbook_path}")
    if not image_directory.is_dir():
        raise FileNotFoundError(f"图片目录不存在: {image_directory}")

    sheet_name, rows, workbook_issues = read_workbook(workbook_path)
    issues.extend(workbook_issues)
    images = sorted(
        path for path in image_directory.rglob("*") if path.is_file() and path.suffix.casefold() in IMAGE_EXTENSIONS
    )

    if expected_count is not None and len(rows) != expected_count:
        issues.append(
            _issue("error", "row_count_mismatch", "有效记录数与预期不一致", expected=expected_count, actual=len(rows))
        )
    if expected_count is not None and len(images) != expected_count:
        issues.append(
            _issue("error", "image_count_mismatch", "图片数与预期不一致", expected=expected_count, actual=len(images))
        )

    record_locations: dict[str, list[int]] = defaultdict(list)
    referenced_names: dict[str, list[int]] = defaultdict(list)
    for record in rows:
        row_number = int(record["__row_number"])
        record_code = normalized_name(record.get("入库编号"))
        image_name = normalized_name(record.get("高清原图文件名"))
        if not record_code:
            issues.append(_issue("error", "missing_record_code", "记录缺少入库编号", row=row_number))
        else:
            record_locations[record_code].append(row_number)
        if not image_name:
            issues.append(_issue("error", "missing_image_name", "记录缺少高清原图文件名", row=row_number))
        else:
            referenced_names[image_name].append(row_number)
        if not clean_text(record.get("作品名称")):
            issues.append(_issue("warning", "missing_title", "记录缺少作品名称", row=row_number))
        if not clean_text(record.get("作者姓名")):
            issues.append(_issue("warning", "missing_artist", "记录缺少作者姓名", row=row_number))

    for record_code, row_numbers in sorted(record_locations.items()):
        if len(row_numbers) > 1:
            issues.append(
                _issue("error", "duplicate_record_code", "入库编号重复", record_code=record_code, rows=row_numbers)
            )
    for image_name, row_numbers in sorted(referenced_names.items()):
        if len(row_numbers) > 1:
            issues.append(
                _issue("error", "duplicate_image_reference", "多条记录引用同一图片名", filename=image_name, rows=row_numbers)
            )

    image_lookup: dict[str, list[Path]] = defaultdict(list)
    empty_images: list[str] = []
    for image in images:
        image_lookup[image.name.casefold()].append(image)
        if image.stat().st_size == 0:
            empty_images.append(str(image.relative_to(image_directory)))
    if empty_images:
        issues.append(_issue("error", "empty_images", "存在空图片文件", files=empty_images))

    for key, paths in sorted(image_lookup.items()):
        if len(paths) > 1:
            issues.append(
                _issue(
                    "error",
                    "duplicate_image_filename",
                    "图片文件名大小写归一化后重复",
                    filename=key,
                    files=[str(path.relative_to(image_directory)) for path in paths],
                )
            )

    missing_references = sorted(name for name in referenced_names if name not in image_lookup)
    if missing_references:
        issues.append(
            _issue("error", "missing_images", "主表引用的图片不存在", count=len(missing_references), files=missing_references)
        )

    orphan_images = sorted(
        str(paths[0].relative_to(image_directory)) for name, paths in image_lookup.items() if name not in referenced_names
    )
    if orphan_images:
        issues.append(
            _issue("error", "orphan_images", "图片未被主表引用", count=len(orphan_images), files=orphan_images)
        )

    image_manifest: list[dict[str, Any]] = []
    hash_groups: dict[str, list[str]] = defaultdict(list)
    for image in images:
        digest = sha256_file(image) if hash_images else None
        relative = str(image.relative_to(image_directory)).replace("\\", "/")
        image_manifest.append({"path": relative, "size": image.stat().st_size, "sha256": digest})
        if digest:
            hash_groups[digest].append(relative)
    duplicate_hashes = {digest: paths for digest, paths in hash_groups.items() if len(paths) > 1}
    for digest, paths in sorted(duplicate_hashes.items()):
        issues.append(
            _issue("warning", "duplicate_image_content", "多张图片内容完全相同", sha256=digest, files=paths)
        )

    manifest_digest = hashlib.sha256()
    manifest_digest.update(sha256_file(workbook_path).encode("ascii"))
    for entry in image_manifest:
        manifest_digest.update(json.dumps(entry, ensure_ascii=False, sort_keys=True).encode("utf-8"))

    errors = sum(issue["severity"] == "error" for issue in issues)
    warnings = sum(issue["severity"] == "warning" for issue in issues)
    return {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "valid": errors == 0,
        "workbook": str(workbook_path),
        "workbook_sha256": sha256_file(workbook_path),
        "sheet": sheet_name,
        "image_directory": str(image_directory),
        "expected_count": expected_count,
        "row_count": len(rows),
        "image_count": len(images),
        "referenced_image_count": len(referenced_names),
        "manifest_sha256": manifest_digest.hexdigest(),
        "error_count": errors,
        "warning_count": warnings,
        "issues": issues,
        "images": image_manifest,
    }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="导入前校验国画数据批次")
    parser.add_argument("--workbook", type=Path, required=True, help="规范化 XLSX 主表")
    parser.add_argument("--image-dir", type=Path, required=True, help="图片根目录")
    parser.add_argument("--expected-count", type=int, help="预期作品和图片数量")
    parser.add_argument("--report", type=Path, required=True, help="JSON 校验报告输出路径")
    parser.add_argument("--skip-image-hash", action="store_true", help="跳过图片 SHA-256（仅用于快速预检）")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        report = validate_batch(args.workbook, args.image_dir, args.expected_count, not args.skip_image_hash)
    except Exception as exc:
        print(f"批次校验失败: {exc}", file=sys.stderr)
        return 2
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(
        f"批次校验完成: valid={str(report['valid']).lower()} rows={report['row_count']} "
        f"images={report['image_count']} errors={report['error_count']} warnings={report['warning_count']}"
    )
    print(f"报告: {args.report.resolve()}")
    return 0 if report["valid"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
